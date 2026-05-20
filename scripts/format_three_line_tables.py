#!/usr/bin/env python3
"""将论文 docx 中全部表格改为标准三线表格式。"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path("/Users/ik/Desktop/01_开发项目/zq")
DEFAULT_THESIS = ROOT / "论文-球类比赛成绩统计与显示管理系统(4).docx"

# Word 边框 sz 单位为 1/8 磅：12=1.5pt（顶/底），6=0.75pt（表头下）
THICK = "12"
THIN = "6"


def _border_el(name: str, val: str, sz: str | None = None, color: str = "000000") -> OxmlElement:
    el = OxmlElement(f"w:{name}")
    el.set(qn("w:val"), val)
    if val not in ("nil", "none"):
        if sz:
            el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
    return el


def _clear_tbl_borders(table) -> None:
    tbl_pr = table._tbl.find(qn("w:tblPr"))
    if tbl_pr is None:
        return
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        return
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = tbl_borders.find(qn(f"w:{side}"))
        if node is not None:
            node.set(qn("w:val"), "nil")


def _set_tc_borders(tc, *, top: str, bottom: str, top_sz: str | None, bottom_sz: str | None) -> None:
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is not None:
        tc_pr.remove(borders)
    borders = OxmlElement("w:tcBorders")
    tc_pr.append(borders)

    for side in ("left", "right", "insideH", "insideV"):
        borders.append(_border_el(side, "nil"))
    borders.append(_border_el("top", top, top_sz))
    borders.append(_border_el("bottom", bottom, bottom_sz))


def apply_three_line_table(table) -> None:
    _clear_tbl_borders(table)
    row_count = len(table.rows)
    if row_count == 0:
        return

    for row_idx, row in enumerate(table.rows):
        is_first = row_idx == 0
        is_last = row_idx == row_count - 1
        seen: set[int] = set()

        for cell in row.cells:
            tc = cell._tc
            key = id(tc)
            if key in seen:
                continue
            seen.add(key)

            if is_first:
                top_val, top_sz = "single", THICK
            else:
                top_val, top_sz = "nil", None

            if is_first and not is_last:
                bottom_val, bottom_sz = "single", THIN
            elif is_last:
                bottom_val, bottom_sz = "single", THICK
            else:
                bottom_val, bottom_sz = "nil", None

            _set_tc_borders(tc, top=top_val, bottom=bottom_val, top_sz=top_sz, bottom_sz=bottom_sz)


def main() -> int:
    thesis = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_THESIS
    if not thesis.exists():
        print(f"文件不存在: {thesis}", file=sys.stderr)
        return 1

    backup = thesis.with_suffix(thesis.suffix + ".bak")
    if not backup.exists():
        shutil.copy2(thesis, backup)
        print(f"已备份: {backup}")

    doc = Document(str(thesis))
    for table in doc.tables:
        apply_three_line_table(table)

    doc.save(str(thesis))
    print(f"已将 {len(doc.tables)} 张表格改为三线表: {thesis}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
