#!/usr/bin/env python3
"""Update thesis: table captions, figure insertions, chapter 5 condense."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path("/Users/ik/Desktop/01_开发项目/zq")
THESIS = ROOT / "论文-球类比赛成绩统计与显示管理系统.docx"
IMAGE_DIR = ROOT / "docs/imge"

TABLE_CAPTIONS = [
    "表3-1 审计日志表结构",
    "表4-1 数据集规模统计",
    "表4-2 实时计分接口说明",
    "表4-3 管理查询接口说明",
    "表4-4 系统功能测试覆盖度",
    "表4-5 实时延迟测试结果",
    "表4-6 吞吐量测试结果",
    "表4-7 多球类适配性测试汇总",
    "表4-8 多终端同步精度测试",
    "表4-9 显示视觉效果评估",
    "表4-10 与人工记录方式对比",
    "表4-11 与STACT系统功能对比",
    "表4-12 综合性能对比",
]

FIGURES = [
    {
        "file": "fig-4-1-manage-overview.png",
        "caption": "图4-1 管理端系统总览界面（brsms-demo）",
        "anchor": "数据集规模统计",
        "ref_text": "如图4-1所示，管理端总览页面展示了比赛、球队、球员与技术统计等核心数据的规模概况。",
        "position": "before",
    },
    {
        "file": "fig-4-2-match-directory.png",
        "caption": "图4-2 比赛目录查询界面（brsms-demo）",
        "anchor": "实时计分接口",
        "ref_text": "如图4-2所示，比赛目录支持按队伍、场馆等条件检索，并提供详情与编辑入口。",
        "position": "before",
    },
    {
        "file": "fig-4-3-match-score-entry.png",
        "caption": "图4-3 比赛成绩录入界面（brsms-demo）",
        "anchor": "管理查询接口",
        "ref_text": "如图4-3所示，录入页面将比赛基础信息与球员技术统计整合在同一表单中完成持久化。",
        "position": "before",
    },
    {
        "file": "fig-4-4-team-archives.png",
        "caption": "图4-4 球队档案管理界面（brsms-demo）",
        "anchor": "管理查询接口（续）",
        "ref_text": "如图4-4所示，球队档案模块支持按球队、城市、教练等维度查询与维护。",
        "position": "before",
    },
    {
        "file": "fig-4-5-standings-table.png",
        "caption": "图4-5 积分榜与统计展示界面（brsms-demo）",
        "anchor": "4.3.1 系统功能测试",
        "ref_text": "如图4-5所示，积分榜页面以表格形式展示各队胜场、负场、积分与胜率等统计指标。",
        "position": "before",
    },
]

CH5_PARAGRAPHS = [
    (
        "本项目围绕基层体育赛事信息化管理的实际需求，设计并实现了球类比赛成绩统计与显示管理系统"
        "（Ball-Game Result Statistics Management System，BRSMS）。系统采用端—边—云分层架构，"
        "在终端侧完成赛事事件的实时采集与本地缓存，在边缘侧实现数据汇聚、协议转换与离线容错，"
        "在云端侧提供成绩管理、统计分析与多屏展示等能力。通过 Web 管理端原型与 brsms-demo 演示系统，"
        "项目完成了实时计分、比赛目录查询、成绩录入、球队档案维护、积分榜统计等核心功能的开发与验证，"
        "有效缓解了传统赛事管理中信息分散、同步滞后、多项目规则难以统一配置等问题，"
        "为高校体育教学、社区联赛及综合运动会等场景提供了可落地的信息化参考方案。"
    ),
    (
        "在技术创新方面，本项目主要形成了以下几方面成果。其一，提出终端—边缘—云端三层离线容错架构，"
        "各层均具备本地持久化与断网续传能力，确保弱网或断网环境下赛事数据不丢失、业务不中断。"
        "其二，设计 BGEMM（Ball-General Event Meta-Model）球类通用事件元模型，以六元组结构统一抽象"
        "篮球、排球、羽毛球、乒乓球等多种球类的事件特征，并结合 Lua 脚本热插拔机制实现规则秒级切换，"
        "显著提升了系统的跨项目适配能力。其三，构建 WebSocket 与 MQTT 双协议融合的实时通信总线，"
        "分别面向浏览器大屏与嵌入式终端，端到端传输延迟控制在 200 ms 以内，满足赛事实时播报需求。"
        "其四，采用 ESP32-S3 等开源硬件方案，将终端设备 BOM 成本控制在 317 元以内，"
        "并开放硬件设计、固件源码与部署镜像，降低了系统推广与二次开发的门槛。"
    ),
    (
        "实验测试进一步验证了系统的可用性与性能表现。功能测试覆盖实时计分、比赛管理、"
        "多球类规则切换、大屏显示等模块，测试通过率 100%；实时性能测试表明，"
        "单客户端计分延迟为 45—110 ms，500 并发连接下 CPU 占用低于 35%，丢包率为零；"
        "多终端同步显示时间差小于 100 ms，视觉效果主观评分达 4.4 分（5 分制）。"
        "与人工记录及 STACT 系统的对比实验表明，本系统在记录效率、数据维度与操作便捷性方面"
        "均具备明显优势，能够支撑基层赛事从赛前准备、赛中执行到赛后统计的全流程管理。"
    ),
    (
        "与此同时，系统仍存在若干有待完善的不足。当前比分与事件录入主要依赖裁判人工操作，"
        "计算机视觉辅助判罚功能尚未实现，对高水平赛事中边界球、触网等精细判罚场景的支撑仍显不足；"
        "云端服务虽已搭建 SpringBoot 基础框架，但高级可视化、细粒度权限控制及与外部教务、"
        "报名系统的深度对接仍需进一步开发；用户界面以 Vue3 Web 应用为主，"
        "虽具备 PWA 离线能力，但在裁判员移动操作场景下，原生 App 的交互体验与离线性能仍有提升空间。"
    ),
    (
        "展望未来，本项目将在现有成果基础上持续迭代优化。在智能化方向，"
        "计划引入 YOLOv8 等目标检测模型，在篮球出界、排球触网、羽毛球落点判定等典型场景开展"
        "自动判罚辅助验证，减轻裁判工作负担、提升判罚一致性。在功能扩展方向，"
        "将进一步完善 BGEMM 元模型，支持足球、网球等更多运动项目，并建设规则脚本共享社区，"
        "形成开放生态。在应用体验方向，将基于 Flutter 或 React Native 开发跨平台移动端应用，"
        "为裁判、技术官员与观众提供差异化功能；同时探索数字孪生与虚拟场馆展示，"
        "拓展赛事传播的沉浸式体验。通过上述工作，BRSMS 有望从演示原型逐步演进为"
        "可在更大范围基层赛事中稳定部署的通用成绩管理平台。"
    ),
]


def insert_paragraph_before(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def set_caption_style(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.size = Pt(10.5)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_caption_before_table(table, caption: str) -> None:
    tbl = table._tbl
    prev = tbl.getprevious()
    if prev is not None and prev.tag == qn("w:p"):
        para = Paragraph(prev, table._parent)
        if para.text.strip() == caption:
            set_caption_style(para)
            return
        if para.text.strip().startswith(("表3-", "表4-")):
            para.clear()
            run = para.add_run(caption)
            set_caption_style(para)
            return
    new_p = OxmlElement("w:p")
    tbl.addprevious(new_p)
    cap = Paragraph(new_p, table._parent)
    cap.add_run(caption)
    set_caption_style(cap)


def find_paragraph(doc: Document, anchor: str) -> Paragraph:
    for p in doc.paragraphs:
        if anchor in p.text:
            return p
    raise ValueError(f"Anchor not found: {anchor}")


def insert_figure_block(doc: Document, fig: dict) -> None:
    img_path = IMAGE_DIR / fig["file"]
    if not img_path.exists():
        raise FileNotFoundError(img_path)

    anchor_para = find_paragraph(doc, fig["anchor"])

    with Image.open(img_path) as im:
        w, h = im.size
    max_w_cm = 14.0
    ratio = h / w
    width = Cm(max_w_cm)
    height = Cm(max_w_cm * ratio)

    cap_para = insert_paragraph_before(anchor_para, fig["caption"])
    set_caption_style(cap_para)

    img_para = insert_paragraph_before(cap_para)
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_para.add_run()
    run.add_picture(str(img_path), width=width, height=height)

    insert_paragraph_before(img_para, fig["ref_text"])


def condense_chapter5(doc: Document) -> None:
    start_idx = None
    end_idx = None
    for i, p in enumerate(doc.paragraphs):
        if "第五章 总结与展望" in p.text:
            start_idx = i
        if start_idx is not None and "参考文献" in p.text:
            end_idx = i
            break

    if start_idx is None or end_idx is None:
        raise ValueError("Chapter 5 boundaries not found")

    title_para = doc.paragraphs[start_idx]

    to_remove = []
    for i in range(start_idx + 1, end_idx):
        to_remove.append(doc.paragraphs[i]._element)

    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)

    prev = title_para
    for text in CH5_PARAGRAPHS:
        para = insert_paragraph_after(prev, text)
        para.style = "Normal"
        prev = para


def main() -> int:
    if not THESIS.exists():
        print(f"Missing thesis: {THESIS}", file=sys.stderr)
        return 1

    doc = Document(str(THESIS))

    if len(doc.tables) != len(TABLE_CAPTIONS):
        print(f"Warning: {len(doc.tables)} tables vs {len(TABLE_CAPTIONS)} captions", file=sys.stderr)

    for table, caption in zip(doc.tables, TABLE_CAPTIONS):
        add_caption_before_table(table, caption)

    for fig in reversed(FIGURES):
        insert_figure_block(doc, fig)

    condense_chapter5(doc)

    doc.save(str(THESIS))
    print(f"Updated: {THESIS}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
