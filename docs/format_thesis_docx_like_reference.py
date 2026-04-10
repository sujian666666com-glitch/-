from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SRC = Path("/Users/ik/Desktop/01_开发项目/zq/docs/论文-球类比赛成绩统计与显示管理系统.docx")
REF = Path("/Users/ik/Desktop/02_学习与资料/毕业设计/Design/正文/毕业论文_完整版_规范化.docx")
OUT = Path("/Users/ik/Desktop/01_开发项目/zq/docs/论文-球类比赛成绩统计与显示管理系统-规范化.docx")


RE_ABSTRACT = re.compile(r"^摘\s*要$")
RE_TOC = re.compile(r"^目\s*录$")
RE_KEYWORDS = re.compile(r"^关键词[：:]")
RE_CHAPTER = re.compile(r"^第\s*[一二三四五六七八九十百0-9]+\s*章")
RE_SEC_2 = re.compile(r"^\d+\.\d+\s+")
RE_SEC_3 = re.compile(r"^\d+\.\d+\.\d+\s+")
RE_SUBTITLE = re.compile(r"^[一二三四五六七八九十]+[、.]")
RE_REFERENCES = re.compile(r"^参考文献$")
RE_ACK = re.compile(r"^致谢$")


def set_east_asia_font(run, east_asia: str) -> None:
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")


def set_run_font(run, east_asia: str, size_pt: float, bold: bool | None = None, underline: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    set_east_asia_font(run, east_asia)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if underline is not None:
        run.underline = underline


def clear_paragraph_indents(paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.left_indent = None
    fmt.right_indent = None
    fmt.first_line_indent = None


def format_cover_paragraph(paragraph, idx: int) -> None:
    text = paragraph.text.strip()
    fmt = paragraph.paragraph_format
    clear_paragraph_indents(paragraph)
    fmt.space_after = Pt(0)

    if idx == 0:
        fmt.space_before = Pt(100)
        return
    if idx == 1 and text:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, "黑体", 26, True)
        return
    if idx == 2:
        fmt.space_before = Pt(20)
        return
    if idx == 3 and text:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, "黑体", 26, True)
        return
    if idx == 4:
        fmt.space_before = Pt(75)
        return
    if idx == 5 and text:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fmt.space_before = Pt(7)
        for run in paragraph.runs:
            set_run_font(run, "黑体", 21.5, True, True)
        return
    if idx in {6, 11}:
        fmt.space_before = Pt(75 if idx == 11 else 100)
        return
    if idx in {7, 8, 9, 10} and text:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Pt(94)
        fmt.space_before = Pt(11)
        for run in paragraph.runs:
            set_run_font(run, "宋体", 15, True)
        return
    if idx == 12 and text:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, "宋体", 14, False)


def format_heading(paragraph, size_pt: float, align: WD_ALIGN_PARAGRAPH, before_pt: float) -> None:
    fmt = paragraph.paragraph_format
    clear_paragraph_indents(paragraph)
    paragraph.alignment = align
    fmt.space_before = Pt(before_pt)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for run in paragraph.runs:
        set_run_font(run, "宋体", size_pt, True)


def format_body_paragraph(paragraph) -> None:
    fmt = paragraph.paragraph_format
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.left_indent = None
    fmt.right_indent = None
    fmt.first_line_indent = Pt(24)
    fmt.space_before = Pt(3)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.18
    for run in paragraph.runs:
        size = run.font.size.pt if run.font.size else 12
        bold = bool(run.bold) if run.bold is not None else False
        set_run_font(run, "宋体", size, bold)


def format_special_label_paragraph(paragraph) -> None:
    fmt = paragraph.paragraph_format
    clear_paragraph_indents(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt.space_before = Pt(3)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.18
    for i, run in enumerate(paragraph.runs):
        set_run_font(run, "宋体", 12, bool(run.bold) if run.bold is not None else i == 0)


def format_toc_paragraph(paragraph, level: int | None) -> None:
    fmt = paragraph.paragraph_format
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    fmt.first_line_indent = None
    fmt.left_indent = Pt(0 if level == 1 else 17 if level == 2 else 34 if level == 3 else 0)
    for run in paragraph.runs:
        set_run_font(run, "宋体", 10.5, bool(run.bold) if run.bold is not None else False)


def infer_toc_level(text: str) -> int | None:
    if RE_CHAPTER.match(text):
        return 1
    if re.match(r"^\d+\.\d+\s+", text):
        return 2
    if re.match(r"^\d+\.\d+\.\d+\s+", text):
        return 3
    return None


def remove_all_headers_footers(doc: Document) -> None:
    for section in doc.sections:
        sect_pr = section._sectPr
        for tag in ("w:headerReference", "w:footerReference"):
            for node in list(sect_pr.findall(qn(tag))):
                sect_pr.remove(node)


def apply_section_layout(doc: Document, ref_doc: Document) -> None:
    ref_section = ref_doc.sections[0]
    for section in doc.sections:
        section.page_width = ref_section.page_width
        section.page_height = ref_section.page_height
        section.top_margin = ref_section.top_margin
        section.left_margin = ref_section.left_margin
        section.right_margin = ref_section.right_margin
        section.header_distance = ref_section.header_distance
        section.footer_distance = ref_section.footer_distance
        # 参考文档底边距为 0cm，直接复制容易造成分页和页脚异常，这里保留较安全的 2.54cm。
        section.bottom_margin = Cm(2.54)
        section.different_first_page_header_footer = False
        section.start_type = WD_SECTION.NEW_PAGE if section is not doc.sections[0] else section.start_type


def normalize_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                fmt = paragraph.paragraph_format
                fmt.first_line_indent = None
                fmt.left_indent = None
                fmt.right_indent = None
                fmt.space_before = Pt(0)
                fmt.space_after = Pt(0)
                fmt.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(run, "宋体", 10.5, row_idx == 0)


def main() -> None:
    doc = Document(str(SRC))
    ref_doc = Document(str(REF))
    apply_section_layout(doc, ref_doc)
    remove_all_headers_footers(doc)

    in_cover = True
    in_abstract = False
    in_toc = False
    body_started = False

    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()

        if in_cover:
            if RE_ABSTRACT.match(text):
                in_cover = False
                in_abstract = True
                format_heading(paragraph, 15.5, WD_ALIGN_PARAGRAPH.CENTER, 12)
            else:
                format_cover_paragraph(paragraph, idx)
            continue

        if RE_TOC.match(text):
            in_abstract = False
            in_toc = True
            format_heading(paragraph, 15.5, WD_ALIGN_PARAGRAPH.CENTER, 12)
            continue

        if in_abstract:
            if not text:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                continue
            if RE_KEYWORDS.match(text):
                format_special_label_paragraph(paragraph)
            else:
                format_body_paragraph(paragraph)
            continue

        if in_toc and RE_CHAPTER.match(text):
            in_toc = False
            body_started = True

        if in_toc:
            if not text:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                continue
            level = infer_toc_level(text)
            format_toc_paragraph(paragraph, level)
            continue

        if not text:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            continue

        if RE_REFERENCES.match(text) or RE_ACK.match(text) or RE_CHAPTER.match(text):
            format_heading(paragraph, 15.5, WD_ALIGN_PARAGRAPH.CENTER, 12)
            body_started = True
            continue
        if RE_SEC_3.match(text):
            format_heading(paragraph, 14, WD_ALIGN_PARAGRAPH.LEFT, 10)
            continue
        if RE_SEC_2.match(text):
            format_heading(paragraph, 15, WD_ALIGN_PARAGRAPH.LEFT, 12)
            continue
        if RE_SUBTITLE.match(text):
            format_heading(paragraph, 12, WD_ALIGN_PARAGRAPH.LEFT, 3)
            continue

        if body_started:
            format_body_paragraph(paragraph)

    for table in doc.tables:
        normalize_table(table)

    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    doc.save(str(OUT))
    print(str(OUT))


if __name__ == "__main__":
    main()
